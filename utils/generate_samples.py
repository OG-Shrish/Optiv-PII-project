"""Utility script to generate realistic synthetic business evidence documents for testing."""
from pathlib import Path
import pymupdf as fitz
import docx
from pptx import Presentation
from pptx.util import Inches, Pt


def generate_all_samples(target_dir: Path):
    """Generate sample PDF, DOCX, PPTX, and TXT files containing realistic business PII."""
    target_dir.mkdir(parents=True, exist_ok=True)

    # 1. Plain Text Sample
    txt_path = target_dir / "Customer_Support_Evidence.txt"
    txt_content = (
        "CONFIDENTIAL AUDIT EVIDENCE - TICKET DISPATCH LOG\n\n"
        "Incident ID: INC-84920\n"
        "Date of Event: October 12, 2026\n"
        "Customer Name: John Smith\n"
        "Customer Email: john.smith@example.com\n"
        "Customer Phone: +91 9876543210\n"
        "Secondary Contact: John Smith\n"
        "Account Owner: Robert Vance\n"
        "Assigned Investigator: EMP10293\n\n"
        "Summary Description:\n"
        "Customer John Smith reported unauthorized billing activity. The card ending was verified.\n"
        "Credit Card on file: 4111 1111 1111 1111.\n"
        "Customer Government ID: GVT-123456.\n"
        "Originating Gateway IP Address: 192.168.1.10.\n"
        "System latency was 42ms with 99.9% uptime across 12 cluster nodes.\n\n"
        "Resolution Notes:\n"
        "Employee ID: EMP-10293 cleared the fraudulent transaction.\n"
        "Confirmation emailed to john.smith@example.com."
    )
    txt_path.write_text(txt_content, encoding="utf-8")

    # 2. DOCX Sample
    docx_path = target_dir / "HR_Incident_Report.docx"
    doc = docx.Document()
    doc.add_heading("Internal Security Investigation Report", level=1)
    
    p1 = doc.add_paragraph("Classification: RESTRICTED INTERNAL REVIEW")
    p1.bold = True
    
    doc.add_heading("1. Subject Identification", level=2)
    doc.add_paragraph(
        "During internal audit review, evidence was collected regarding department access anomalies.\n"
        "Primary Employee Name: Sarah Connor\n"
        "Employee ID: EMP88341\n"
        "Corporate Email: sarah.connor@cyberdyne.org\n"
        "Contact Phone: (212) 555-1234"
    )

    doc.add_heading("2. Witness Statements & Roster", level=2)
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Employee"
    table.cell(0, 1).text = "Staff ID"
    table.cell(0, 2).text = "Contact Email"
    
    table.cell(1, 0).text = "Michael Scott"
    table.cell(1, 1).text = "EMP10293"
    table.cell(1, 2).text = "mscott@dunder.com"

    table.cell(2, 0).text = "Alice Henderson"
    table.cell(2, 1).text = "EMP44912"
    table.cell(2, 2).text = "a.henderson@securecorp.net"

    doc.add_paragraph(
        "Quarterly financial volume totaled 2,450,000 USD. No unauthorized fund transfer occurred."
    )
    doc.save(docx_path)

    # 3. PPTX Sample
    pptx_path = target_dir / "Cadence_Financial_Group_Org_Pack.pptx"
    prs = Presentation()
    
    # Slide 1: Title
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    tb1 = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    tf1 = tb1.text_frame
    p_title = tf1.paragraphs[0]
    p_title.text = "Cadence Financial Group - Global Leadership"
    p_title.font.size = Pt(28)
    p_title.font.bold = True

    # Slide 2: Operations Contact Pack
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.4), Inches(5))
    tf2 = tb2.text_frame
    p_h = tf2.paragraphs[0]
    p_h.text = "Operations Incident Escalation Directory"
    p_h.font.size = Pt(22)
    p_h.font.bold = True

    p_body = tf2.add_paragraph()
    p_body.text = (
        "Incident Lead: David Miller\n"
        "Direct Phone: +1 415 555 2671\n"
        "Escalation Email: d.miller@cadencefinancial.com\n"
        "Employee Number: EMP99420\n"
        "Backup Contact: David Miller\n"
        "Server Subnet Gateway: 10.240.12.1"
    )
    p_body.font.size = Pt(16)
    prs.save(pptx_path)

    # 4. PDF Sample
    pdf_path = target_dir / "Executive_Briefing.pdf"
    pdf_doc = fitz.open()
    
    # Page 1
    page1 = pdf_doc.new_page()
    page1.insert_text(
        (50, 72),
        "CADENCE ENTERPRISE AUDIT BRIEFING\n\n"
        "Document Ref: AUD-2026-99\n"
        "Reviewing Officer: Katherine Warren\n"
        "Email: k.warren@enterprise-sec.io\n"
        "Direct Phone: +91 9876543210\n"
        "Employee ID: EMP55210\n\n"
        "Overview:\n"
        "Total processed transactions reached 54,200 requests during peak load.\n"
        "Government ID verification was completed for ticket GVT-881290.\n"
        "Associated Internal Server IP Address: 172.16.254.1\n\n"
        "Katherine Warren verified compliance with security policy Section 12.",
        fontsize=12
    )

    # Page 2
    page2 = pdf_doc.new_page()
    page2.insert_text(
        (50, 72),
        "APPENDIX A: REGULATORY IDENTIFIERS\n\n"
        "SSN on file for account verification: 987-65-4321\n"
        "Date of Birth: 12/04/1998\n"
        "Account Executive: Katherine Warren\n"
        "All customer records are retained pursuant to 5-year compliance mandate.\n"
        "No unencrypted records are transmitted outside corporate perimeter.",
        fontsize=12
    )
    pdf_doc.save(pdf_path)
    pdf_doc.close()

    print(f"Generated sample evidence files in: {target_dir}")


if __name__ == "__main__":
    generate_all_samples(Path(__file__).resolve().parent.parent / "sample_documents")

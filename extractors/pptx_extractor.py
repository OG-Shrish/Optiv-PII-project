"""PPTX Document Extractor using python-pptx preserving slides, shapes, tables, and notes."""
import io
import logging
from typing import List, Tuple
from pptx import Presentation

from models.schemas import DocumentBlock

logger = logging.getLogger(__name__)


def extract_pptx(
    file_bytes: bytes,
    file_name: str,
    document_id: str
) -> Tuple[List[DocumentBlock], bool]:
    """Extract structured document blocks from a PPTX presentation.

    Returns:
        Tuple of (blocks, requires_ocr=False)
    """
    blocks: List[DocumentBlock] = []
    order_counter = 0

    try:
        prs = Presentation(io.BytesIO(file_bytes))
    except Exception as e:
        logger.error(f"Failed to open PPTX stream for {file_name}: {e}")
        raise ValueError(f"Corrupted or invalid PPTX file: {e}")

    for s_idx, slide in enumerate(prs.slides):
        slide_num = s_idx + 1

        # Process shapes and text boxes
        for sh_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                    text = paragraph.text.strip()
                    if text:
                        blocks.append(
                            DocumentBlock(
                                document_id=document_id,
                                file_name=file_name,
                                file_type="pptx",
                                location=f"slide_{slide_num}_shape_{sh_idx + 1}_p_{p_idx + 1}",
                                block_type="slide_text",
                                text=text,
                                order=order_counter,
                                metadata={
                                    "slide": slide_num,
                                    "shape_id": shape.shape_id,
                                    "paragraph_index": p_idx + 1,
                                }
                            )
                        )
                        order_counter += 1

            # Process tables within slide
            elif shape.has_table:
                table = shape.table
                for r_idx, row in enumerate(table.rows):
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        row_combined = " | ".join(row_cells)
                        blocks.append(
                            DocumentBlock(
                                document_id=document_id,
                                file_name=file_name,
                                file_type="pptx",
                                location=f"slide_{slide_num}_table_{sh_idx + 1}_row_{r_idx + 1}",
                                block_type="slide_table",
                                text=row_combined,
                                order=order_counter,
                                metadata={
                                    "slide": slide_num,
                                    "shape_id": shape.shape_id,
                                    "row_index": r_idx + 1,
                                }
                            )
                        )
                        order_counter += 1

        # Process speaker notes if present
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    blocks.append(
                        DocumentBlock(
                            document_id=document_id,
                            file_name=file_name,
                            file_type="pptx",
                            location=f"slide_{slide_num}_notes",
                            block_type="speaker_notes",
                            text=notes_text,
                            order=order_counter,
                            metadata={"slide": slide_num, "is_notes": True}
                        )
                    )
                    order_counter += 1
        except Exception as notes_err:
            logger.debug(f"Could not read speaker notes for slide {slide_num}: {notes_err}")

    return blocks, False

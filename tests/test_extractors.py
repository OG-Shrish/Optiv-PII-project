"""Tests for document extractors (PDF, DOCX, PPTX, TXT)."""
import io
import pymupdf as fitz
import docx
from pptx import Presentation
from pptx.util import Inches

from extractors import extract_document
from extractors.pdf_extractor import extract_pdf
from extractors.docx_extractor import extract_docx
from extractors.pptx_extractor import extract_pptx
from extractors.txt_extractor import extract_txt


def test_txt_extraction():
    text_content = "First paragraph with normal business text.\n\nSecond paragraph with Contact: Alice Smith."
    file_bytes = text_content.encode("utf-8")
    
    blocks, requires_ocr, status = extract_document(file_bytes, "test_doc.txt", "doc_1")
    
    assert status == "SUCCESS"
    assert requires_ocr is False
    assert len(blocks) == 2
    assert blocks[0].location == "paragraph_1"
    assert "First paragraph" in blocks[0].text
    assert blocks[1].location == "paragraph_2"
    assert "Alice Smith" in blocks[1].text


def test_docx_extraction():
    # Build sample DOCX in memory
    doc = docx.Document()
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("This is the main audit finding paragraph.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Employee"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "John Smith"
    table.cell(1, 1).text = "Auditor"
    
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    blocks, requires_ocr, status = extract_document(docx_bytes, "audit.docx", "doc_2")

    assert status == "SUCCESS"
    assert len(blocks) >= 2
    # Verify structure preserved: headings and table rows captured
    locations = [b.location for b in blocks]
    assert any("heading" in loc for loc in locations)
    assert any("table" in loc for loc in locations)


def test_pptx_extraction():
    # Build sample PPTX in memory
    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    tx_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    tf = tx_box.text_frame
    tf.text = "Cadence Financial Group Organizational Pack"
    p2 = tf.add_paragraph()
    p2.text = "Contact: Sarah Jenkins\nEmail: s.jenkins@example.com"

    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    blocks, requires_ocr, status = extract_document(pptx_bytes, "org_pack.pptx", "doc_3")

    assert status == "SUCCESS"
    assert len(blocks) >= 1
    assert "slide_1" in blocks[0].location
    assert any("Sarah Jenkins" in b.text for b in blocks)


def test_pdf_extraction():
    # Build sample PDF in memory
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Incident Response Plan\nAuthor: Michael Scott\nEmail: mscott@dunder.com")
    pdf_bytes = doc.tobytes()
    doc.close()

    blocks, requires_ocr, status = extract_document(pdf_bytes, "plan.pdf", "doc_4")

    assert status == "SUCCESS"
    assert len(blocks) >= 1
    assert "page_1" in blocks[0].location
    assert any("Incident Response Plan" in b.text for b in blocks)


def test_scanned_pdf_detection():
    # PDF page with empty text should trigger requires_ocr
    doc = fitz.open()
    doc.new_page()  # Blank page
    pdf_bytes = doc.tobytes()
    doc.close()

    blocks, requires_ocr = extract_pdf(pdf_bytes, "scanned.pdf", "doc_5")

    assert requires_ocr is True
    assert len(blocks) == 1
    assert "SCANNED_PAGE_REQUIRES_OCR" in blocks[0].text
